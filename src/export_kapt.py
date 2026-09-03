from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from .clean_kapt import clean_kapt
from .config import ensure_directories, load_settings


COLUMN_LABELS = {
    "kapt_code": "K-apt 단지코드",
    "complex_name": "단지명",
    "sigungu": "시군구",
    "dong": "법정동",
    "jibun": "지번",
    "road_address": "도로명주소",
    "legal_address": "법정동주소",
    "households": "세대수",
    "buildings": "동수",
    "approval_date": "사용승인일",
    "apartment_age": "단지연령(년)",
    "age_group": "연령구간",
    "parking_total": "총주차대수",
    "parking_ground": "지상주차대수",
    "parking_underground": "지하주차대수",
    "parking_per_household": "세대당주차대수",
    "households_per_building": "동당세대수",
    "heating": "난방방식",
    "mixed_use": "단지분류",
    "latitude": "위도",
    "longitude": "경도",
    "is_500plus": "500세대이상",
    "is_1000plus": "1000세대이상",
    "is_under_10years": "10년이하",
    "is_over_20years": "20년초과",
    "is_over_30years": "30년초과",
}


COLUMN_DESCRIPTIONS = {
    "kapt_code": "K-apt에서 부여한 공동주택 단지 식별자",
    "complex_name": "K-apt 등록 단지명",
    "sigungu": "부산광역시 구·군",
    "dong": "법정동 명칭",
    "jibun": "법정동주소에서 추출·정규화한 지번",
    "road_address": "K-apt 도로명주소 원문",
    "legal_address": "K-apt 법정동주소 원문",
    "households": "단지 전체 세대수",
    "buildings": "단지 내 동 수",
    "approval_date": "사용승인일",
    "apartment_age": "기준일 현재 사용승인 후 경과 연수",
    "age_group": "단지연령 구간",
    "parking_total": "지상·지하 주차대수 합계",
    "parking_ground": "지상 주차대수",
    "parking_underground": "지하 주차대수",
    "parking_per_household": "총주차대수 / 세대수",
    "households_per_building": "세대수 / 동수",
    "heating": "난방 방식",
    "mixed_use": "K-apt 단지 분류",
    "latitude": "WGS84 위도",
    "longitude": "WGS84 경도",
    "is_500plus": "500세대 이상 여부",
    "is_1000plus": "1,000세대 이상 여부",
    "is_under_10years": "단지연령 10년 이하 여부",
    "is_over_20years": "단지연령 20년 초과 여부",
    "is_over_30years": "단지연령 30년 초과 여부",
}


def _excel_safe(frame: pd.DataFrame) -> pd.DataFrame:
    safe = frame.copy()
    for column in safe.select_dtypes(include=["object", "string"]).columns:
        safe[column] = safe[column].map(
            lambda value: "'" + value
            if isinstance(value, str) and value.startswith(("=", "+", "-", "@"))
            else value
        )
    return safe


def _quality_summary(cleaned: pd.DataFrame) -> pd.DataFrame:
    checks = [
        ("전체 단지 수", len(cleaned), "행 수"),
        ("고유 K-apt 단지코드", cleaned["kapt_code"].nunique(), "중복 제거 단지 수"),
        ("단지코드 중복", cleaned["kapt_code"].duplicated().sum(), "0이 정상"),
        ("단지명 누락", cleaned["complex_name"].isna().sum(), "원천 데이터 결측"),
        ("도로명주소 누락", cleaned["road_address"].isna().sum(), "원천 데이터 결측"),
        ("법정동주소 누락", cleaned["legal_address"].isna().sum(), "원천 데이터 결측"),
        ("지번 누락", cleaned["jibun"].isna().sum(), "법정동주소에서 추출 불가"),
        ("세대수 누락", cleaned["households"].isna().sum(), "원천 데이터 결측"),
        ("주차대수 누락", cleaned["parking_total"].isna().sum(), "원천 데이터 결측"),
    ]
    return pd.DataFrame(checks, columns=["검사항목", "건수", "설명"])


def _format_sheet(worksheet: Any, *, freeze: str = "A2", max_width: int = 45) -> None:
    worksheet.freeze_panes = freeze
    worksheet.auto_filter.ref = worksheet.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 24
    for column_cells in worksheet.columns:
        values = ["" if cell.value is None else str(cell.value) for cell in column_cells[:200]]
        width = min(max(max((len(value) for value in values), default=0) + 2, 10), max_width)
        worksheet.column_dimensions[column_cells[0].column_letter].width = width


def write_kapt_excel(raw: pd.DataFrame, target: Path) -> Path:
    if raw.empty:
        raise RuntimeError("엑셀로 저장할 K-apt 단지 데이터가 없습니다.")
    cleaned = clean_kapt(raw).sort_values(
        ["sigungu", "dong", "complex_name"], na_position="last"
    ).reset_index(drop=True)
    analysis_columns = [column for column in COLUMN_LABELS if column in cleaned]
    analysis = _excel_safe(cleaned[analysis_columns].rename(columns=COLUMN_LABELS))
    raw_export = _excel_safe(raw.copy())
    quality = _quality_summary(cleaned)
    dictionary = pd.DataFrame([
        {
            "분석컬럼": column,
            "엑셀컬럼명": COLUMN_LABELS[column],
            "설명": COLUMN_DESCRIPTIONS.get(column, ""),
        }
        for column in analysis_columns
    ])

    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_name(f"{target.stem}.tmp{target.suffix}")
    with pd.ExcelWriter(temporary, engine="openpyxl", datetime_format="yyyy-mm-dd") as writer:
        analysis.to_excel(writer, sheet_name="분석용_단지목록", index=False)
        raw_export.to_excel(writer, sheet_name="원본_API", index=False)
        quality.to_excel(writer, sheet_name="품질요약", index=False)
        dictionary.to_excel(writer, sheet_name="데이터사전", index=False)
        for worksheet in writer.book.worksheets:
            _format_sheet(worksheet)
        analysis_sheet = writer.book["분석용_단지목록"]
        header_map = {cell.value: cell.column for cell in analysis_sheet[1]}
        for label in ["세대당주차대수", "동당세대수", "위도", "경도"]:
            if label in header_map:
                for cell in analysis_sheet.iter_cols(
                    min_col=header_map[label], max_col=header_map[label], min_row=2
                ):
                    for item in cell:
                        item.number_format = "0.00"

    temporary.replace(target)
    return target


def export_kapt_excel(source: Path | None = None, target: Path | None = None) -> dict[str, int | str]:
    settings = load_settings()
    ensure_directories()
    source = source or Path(settings["paths"]["raw"]) / "kapt" / "busan_complexes.parquet"
    target = target or Path(settings["paths"]["processed"]) / "busan_kapt_apartment_analysis.xlsx"
    if not source.exists():
        raise RuntimeError(f"K-apt 원본 파일이 없습니다: {source}")
    raw = pd.read_parquet(source)
    output = write_kapt_excel(raw, target)
    return {"rows": len(raw), "excel_path": str(output)}
