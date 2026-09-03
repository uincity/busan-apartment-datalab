from pathlib import Path

from openpyxl import load_workbook
import pandas as pd

from src.export_kapt import write_kapt_excel


def test_write_kapt_excel_contains_analysis_raw_quality_and_dictionary_sheets():
    raw = pd.DataFrame([{
        "kaptCode": "K1",
        "kaptName": "테스트아파트",
        "kaptAddr": "부산광역시 중구 영주동 92 테스트아파트",
        "doroJuso": "부산광역시 중구 영주로 73",
        "bjdName": "영주동",
        "as2": "중구",
        "kaptdaCnt": 100,
        "kaptDongCnt": 2,
        "kaptUsedate": "2020-01-01",
    }])
    target = Path(__file__).with_name("_test_kapt_export.xlsx")

    try:
        write_kapt_excel(raw, target)

        workbook = load_workbook(target, read_only=True)
        assert workbook.sheetnames == ["분석용_단지목록", "원본_API", "품질요약", "데이터사전"]
        assert workbook["분석용_단지목록"].max_row == 2
        assert workbook["원본_API"].max_row == 2
        workbook.close()
    finally:
        target.unlink(missing_ok=True)
